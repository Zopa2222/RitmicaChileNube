import openpyxl
import json
from collections import defaultdict

# Cargar el archivo Excel
wb = openpyxl.load_workbook('CENTRO.xlsx')

# Analizar todas las hojas
result = {
    'total_sheets': len(wb.sheetnames),
    'sheet_names': wb.sheetnames,
    'categories_found': set(),
    'sample_data': {}
}

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    
    # Buscar la fila de encabezados (N, NOMBRE, CLUB, CATEGORIA)
    header_row = None
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), 1):
        if 'NOMBRE' in [str(cell).upper() if cell else '' for cell in row]:
            header_row = row_idx
            break
    
    if header_row:
        # Extraer datos de gimnastas
        gymnasts = []
        for row in ws.iter_rows(min_row=header_row + 1, max_row=min(header_row + 50, ws.max_row), values_only=True):
            # Buscar en BANCA A (columnas 0-3)
            if row[1] and row[3]:  # NOMBRE y CATEGORIA
                gymnasts.append({
                    'nombre': str(row[1]),
                    'club': str(row[2]) if row[2] else '',
                    'categoria': str(row[3])
                })
                result['categories_found'].add(str(row[3]))
            
            # Buscar en BANCA B (columnas 4-7)
            if len(row) > 5 and row[5] and len(row) > 7 and row[7]:  # NOMBRE y CATEGORIA
                gymnasts.append({
                    'nombre': str(row[5]),
                    'club': str(row[6]) if row[6] else '',
                    'categoria': str(row[7])
                })
                result['categories_found'].add(str(row[7]))
        
        result['sample_data'][sheet_name] = gymnasts[:10]  # Primeras 10 gimnastas

# Convertir set a lista para JSON
result['categories_found'] = sorted(list(result['categories_found']))

print(json.dumps(result, indent=2, ensure_ascii=False))
