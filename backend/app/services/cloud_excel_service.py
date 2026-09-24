import re
import unicodedata
from collections import OrderedDict
from io import BytesIO

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException


MARKER_KEYWORDS = {
    'PREMIACION': 'PREMIACION',
    'ALMUERZO': 'ALMUERZO',
    'JUECES': 'JUECES',
    'BREAK': 'BREAK',
    'PAUSA': 'PAUSA',
}


class ExcelImportError(ValueError):
    pass


def normalize_text(value):
    text = unicodedata.normalize('NFKD', str(value or ''))
    text = ''.join(char for char in text if not unicodedata.combining(char))
    return re.sub(r'\s+', ' ', text).strip().upper()


def _header_groups(values):
    normalized = [normalize_text(value) for value in values]
    name_positions = [
        index for index, value in enumerate(normalized) if value == 'NOMBRE'
    ]
    category_positions = [
        index
        for index, value in enumerate(normalized)
        if value in {'CATEGORIA', 'CATEGORY'}
    ]
    if not name_positions or not category_positions:
        return []

    club_positions = [
        index for index, value in enumerate(normalized) if value == 'CLUB'
    ]
    groups = []
    for group_index, name_position in enumerate(name_positions):
        next_name = (
            name_positions[group_index + 1]
            if group_index + 1 < len(name_positions)
            else len(values)
        )
        categories = [
            position
            for position in category_positions
            if name_position < position < next_name
        ]
        if not categories:
            continue
        category_position = min(categories)
        clubs = [
            position
            for position in club_positions
            if name_position < position < category_position
        ]
        groups.append({
            'name': name_position,
            'club': min(clubs) if clubs else None,
            'category': category_position,
        })
    return groups


def _marker_from_row(row_number, values):
    non_empty = [
        str(value).strip()
        for value in values
        if value is not None and str(value).strip()
    ]
    if not non_empty:
        return None

    text = ' | '.join(non_empty)
    normalized = normalize_text(text)
    kinds = [
        kind
        for keyword, kind in MARKER_KEYWORDS.items()
        if keyword in normalized
    ]
    if not kinds:
        return None

    duration = None
    match = re.search(
        r'(\d+)\s*(HORA|HORAS|HR|HRS|MIN|MINUTO|MINUTOS)',
        normalized,
    )
    if match:
        amount = int(match.group(1))
        unit = match.group(2)
        duration = amount * 60 if unit.startswith('H') else amount

    primary_candidate = (
        'ALMUERZO' in kinds
        or 'JUECES' in kinds
        or 'PREMIACION' in kinds
        or 'PRINCIPAL' in normalized
        or (duration is not None and duration >= 30)
    )
    return {
        'row': row_number,
        'text': text,
        'kinds': kinds,
        'duration_minutes': duration,
        'primary_candidate': primary_candidate,
    }


def analyze_excel(contents):
    try:
        workbook = openpyxl.load_workbook(
            BytesIO(contents),
            read_only=True,
            data_only=True,
        )
    except (InvalidFileException, OSError, ValueError, KeyError) as error:
        raise ExcelImportError(
            'El archivo no es un Excel .xlsx válido'
        ) from error

    if not workbook.sheetnames:
        raise ExcelImportError('El archivo no contiene hojas')

    sheets = []
    day_worksheets = []
    judge_worksheet = None
    for worksheet in workbook.worksheets:
        if normalize_text(worksheet.title) == 'JUECES':
            judge_worksheet = worksheet
        else:
            day_worksheets.append(worksheet)

    if not day_worksheets:
        workbook.close()
        raise ExcelImportError('Falta la hoja con el orden de paso del día')
    judges = _analyze_judges(judge_worksheet) if judge_worksheet else []
    has_judges_sheet = judge_worksheet is not None

    for sequence, worksheet in enumerate(day_worksheets, 1):
        active_groups = []
        imported_rows = []
        markers = []
        max_content_row = 0

        for row_number, row in enumerate(
            worksheet.iter_rows(values_only=True),
            1,
        ):
            values = list(row)
            if any(
                value is not None and str(value).strip()
                for value in values
            ):
                max_content_row = row_number

            marker = _marker_from_row(row_number, values)
            if marker:
                markers.append(marker)

            detected_groups = _header_groups(values)
            if detected_groups:
                if len(detected_groups) > 2:
                    raise ExcelImportError(
                        f'La hoja {worksheet.title} contiene más de dos bancas'
                    )
                active_groups = detected_groups
                continue
            if not active_groups:
                continue

            for group_index, group in enumerate(active_groups):
                name_value = (
                    values[group['name']]
                    if group['name'] < len(values)
                    else None
                )
                category_value = (
                    values[group['category']]
                    if group['category'] < len(values)
                    else None
                )
                if not name_value or not category_value:
                    continue

                full_name = str(name_value).strip()
                category = str(category_value).strip()
                if (
                    normalize_text(full_name) in {'NOMBRE', 'NAME'}
                    or normalize_text(full_name).startswith('NOMBRE GIMNASTA')
                    or normalize_text(category) in {'CATEGORIA', 'CATEGORY'}
                ):
                    continue

                club = ''
                if group['club'] is not None and group['club'] < len(values):
                    club_value = values[group['club']]
                    club = str(club_value).strip() if club_value else ''

                imported_rows.append({
                    'source_row': row_number,
                    'bench': 'A' if group_index == 0 else 'B',
                    'full_name': full_name,
                    'club_name': club,
                    'category_name': category,
                })

        if not imported_rows:
            raise ExcelImportError(
                f'No se encontraron gimnastas en la hoja {worksheet.title}'
            )

        data_rows = [row['source_row'] for row in imported_rows]
        first_data_row = min(data_rows)
        last_data_row = max(data_rows)
        for marker in markers:
            marker['eligible_cutoff'] = (
                first_data_row < marker['row'] < last_data_row
            )
            if not marker['eligible_cutoff']:
                marker['primary_candidate'] = False

        detected_marker = next(
            (
                marker
                for marker in markers
                if marker['eligible_cutoff']
                and marker['primary_candidate']
            ),
            None,
        )
        sheets.append({
            'sequence': sequence,
            'name': worksheet.title,
            'max_content_row': max_content_row,
            'rows': imported_rows,
            'markers': markers,
            'detected_cutoff_row': (
                detected_marker['row'] if detected_marker else None
            ),
        })

    workbook.close()
    return {
        'version': 1,
        'sheets': sheets,
        'judges': judges,
        'has_judges_sheet': has_judges_sheet,
    }


def _analyze_judges(worksheet):
    """Read judge blocks organized by AM/PM and bench A/B."""
    judges = []
    session = None
    for row_number, row in enumerate(
        worksheet.iter_rows(values_only=True),
        1,
    ):
        values = list(row)
        normalized = [normalize_text(value) for value in values]
        joined = ' '.join(value for value in normalized if value)
        if 'JORNADA AM' in joined:
            session = 'AM'
            continue
        if 'JORNADA PM' in joined:
            session = 'PM'
            continue
        if session is None:
            continue

        if any(value.startswith('BANCA') for value in normalized):
            continue
        is_header = all(
            normalized[index] in {'NOMBRE', 'RUT', 'AREA'}
            for index in (0, 1, 2, 3, 4, 5)
        )
        if is_header:
            continue
        for bench_index, (name_pos, rut_pos, role_pos) in enumerate(((0, 1, 2), (3, 4, 5))):
            name = str(values[name_pos] or '').strip()
            rut = str(values[rut_pos] or '').strip()
            role_text = normalize_text(values[role_pos])
            if not name and not rut and not role_text:
                continue
            normalized_name = normalize_text(name)
            if normalized_name.startswith('NOMBRE Y APELLIDO'):
                continue
            match = re.fullmatch(r'(DA|DB|A|E|L|P)\s*\d*', role_text)
            if not name or not rut or not match:
                raise ExcelImportError(f'Asignación de juez incompleta o rol inválido en Jueces, fila {row_number}')
            judges.append({
                'first_name': name.split()[0],
                'last_name': ' '.join(name.split()[1:]) or name.split()[0],
                'full_name': name,
                'rut': rut,
                'role': match.group(1),
                'bench': 'A' if bench_index == 0 else 'B',
                'session': session,
                'source_row': row_number,
            })

    unique = {}
    for judge in judges:
        key = (judge['session'], judge['bench'], normalize_text(judge['rut']), judge['role'])
        unique[key] = judge
    counts = {}
    for judge in unique.values():
        scope = (judge['session'], judge['bench'], judge['role'])
        counts[scope] = counts.get(scope, 0) + 1
        if counts[scope] > 4:
            raise ExcelImportError(
                f"Máximo 4 jueces {judge['role']} por banca y jornada: "
                f"banca {judge['bench']}, {judge['session']}"
            )
    return list(unique.values())


def build_sheet_plan(sheet, cutoff_row):
    grouped = OrderedDict()
    location_by_name = {}
    for row in sheet['rows']:
        session = 'AM' if cutoff_row and row['source_row'] < cutoff_row else (
            'PM' if cutoff_row else None
        )
        location = (
            sheet['sequence'],
            row['bench'],
            session,
        )
        previous_location = location_by_name.get(row['category_name'])
        if previous_location and previous_location != location:
            raise ExcelImportError(
                f'La categoría {row["category_name"]} aparece en más de una '
                f'banca o jornada en la hoja {sheet["name"]}'
            )
        location_by_name[row['category_name']] = location

        key = (row['bench'], session, row['category_name'])
        category = grouped.setdefault(
            key,
            {
                'name': row['category_name'],
                'bench': row['bench'],
                'session': session,
                'first_row': row['source_row'],
                'last_row': row['source_row'],
                'gymnasts': [],
            },
        )
        category['last_row'] = row['source_row']
        category['gymnasts'].append({
            'full_name': row['full_name'],
            'club_name': row['club_name'],
            'source_row': row['source_row'],
            'passing_order': len(category['gymnasts']),
        })

    order_by_scope = {}
    categories = []
    for category in grouped.values():
        scope = (category['bench'], category['session'])
        category['passing_order'] = order_by_scope.get(scope, 0)
        order_by_scope[scope] = category['passing_order'] + 1
        category['gymnast_count'] = len(category['gymnasts'])
        categories.append(category)
    return categories


def render_preview(analysis, decisions=None, include_gymnasts=False):
    decisions = decisions or {}
    decision_sheets = decisions.get('sheets', {})
    rendered_sheets = []
    total_categories = 0
    total_gymnasts = 0

    for sheet in analysis['sheets']:
        decision = decision_sheets.get(str(sheet['sequence']), {})
        cutoff_row = decision.get(
            'cutoff_row',
            sheet['detected_cutoff_row'],
        )
        categories = build_sheet_plan(sheet, cutoff_row)
        public_categories = []
        summary = {
            'A': {'AM': {'categories': 0, 'gymnasts': 0},
                  'PM': {'categories': 0, 'gymnasts': 0}},
            'B': {'AM': {'categories': 0, 'gymnasts': 0},
                  'PM': {'categories': 0, 'gymnasts': 0}},
        }

        for category in categories:
            public_category = {
                key: value
                for key, value in category.items()
                if include_gymnasts or key != 'gymnasts'
            }
            public_categories.append(public_category)
            total_categories += 1
            total_gymnasts += category['gymnast_count']
            if category['session']:
                bucket = summary[category['bench']][category['session']]
                bucket['categories'] += 1
                bucket['gymnasts'] += category['gymnast_count']

        rendered_sheets.append({
            'sequence': sheet['sequence'],
            'name': sheet['name'],
            'max_content_row': sheet['max_content_row'],
            'markers': sheet['markers'],
            'detected_cutoff_row': sheet['detected_cutoff_row'],
            'selected_cutoff_row': cutoff_row,
            'cutoff_confirmed': bool(decision.get('confirmed')),
            'requires_manual_decision': (
                sheet['detected_cutoff_row'] is None
            ),
            'summary': summary,
            'categories': public_categories,
        })

    return {
        'version': analysis['version'],
        'total_days': len(rendered_sheets),
        'total_categories': total_categories,
        'total_gymnasts': total_gymnasts,
        'judges': analysis.get('judges', []),
        'has_judges_sheet': analysis.get('has_judges_sheet', False),
        'sheets': rendered_sheets,
    }


def validate_cutoff_decisions(analysis, requested_decisions, current=None):
    current = current or {'sheets': {}}
    result = dict(current)
    result['sheets'] = dict(current.get('sheets', {}))
    sheets_by_sequence = {
        sheet['sequence']: sheet for sheet in analysis['sheets']
    }

    for requested in requested_decisions:
        try:
            sequence = int(requested['sequence'])
            cutoff_row = int(requested['cutoff_row'])
        except (KeyError, TypeError, ValueError):
            raise ExcelImportError(
                'Cada decisión requiere sequence y cutoff_row numéricos'
            ) from None

        sheet = sheets_by_sequence.get(sequence)
        if sheet is None:
            raise ExcelImportError(f'No existe el día {sequence}')
        if cutoff_row < 1 or cutoff_row > sheet['max_content_row'] + 1:
            raise ExcelImportError(
                f'La fila de corte del día {sequence} está fuera del rango'
            )

        # Building the plan detects categories split by an invalid cut.
        build_sheet_plan(sheet, cutoff_row)
        result['sheets'][str(sequence)] = {
            'cutoff_row': cutoff_row,
            'confirmed': bool(requested.get('confirmed', True)),
            'source': requested.get('source', 'MANUAL'),
        }

    return result


def confirmed_import_plan(analysis, decisions):
    decision_sheets = decisions.get('sheets', {})
    plan = []
    category_locations = {}
    for sheet in analysis['sheets']:
        decision = decision_sheets.get(str(sheet['sequence']))
        if not decision or not decision.get('confirmed'):
            raise ExcelImportError(
                f'Falta confirmar el corte AM/PM del día {sheet["sequence"]}'
            )
        cutoff_row = decision.get('cutoff_row')
        categories = build_sheet_plan(sheet, cutoff_row)
        for category in categories:
            location = (
                sheet['sequence'],
                category['bench'],
                category['session'],
            )
            previous = category_locations.get(category['name'])
            if previous and previous != location:
                raise ExcelImportError(
                    f'La categoría {category["name"]} aparece en más de un '
                    'día, banca o jornada'
                )
            category_locations[category['name']] = location

        plan.append({
            'sequence': sheet['sequence'],
            'name': sheet['name'],
            'cutoff_row': cutoff_row,
            'categories': categories,
        })
    return plan


def single_day_analysis(analysis, day_sequence):
    """Restrict a workbook analysis to its sole non-judge sheet."""
    sheets = analysis['sheets']
    if len(sheets) != 1:
        raise ExcelImportError(
            'Cada archivo debe contener exactamente una hoja de orden del día'
        )
    sheet = dict(sheets[0])
    sheet['sequence'] = day_sequence
    return {
        **analysis,
        'sheets': [sheet],
    }


def confirmed_single_day_plan(analysis, cutoff_row):
    if len(analysis['sheets']) != 1:
        raise ExcelImportError(
            'Cada archivo debe contener exactamente una hoja de orden del día'
        )
    sheet = analysis['sheets'][0]
    if not cutoff_row:
        raise ExcelImportError('Debe confirmar el corte entre jornadas AM y PM')
    categories = build_sheet_plan(sheet, cutoff_row)
    for category in categories:
        if category['session'] not in {'AM', 'PM'}:
            raise ExcelImportError('No fue posible identificar la jornada de una categoría')
    return {
        'sequence': sheet['sequence'],
        'name': sheet['name'],
        'cutoff_row': cutoff_row,
        'categories': categories,
        'judges': analysis.get('judges', []),
    }
