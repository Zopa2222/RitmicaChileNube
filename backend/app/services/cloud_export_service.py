"""Result exports from the normalized cloud schema."""

from io import BytesIO
from secrets import token_urlsafe
from xml.sax.saxutils import escape

from reportlab.lib.pdfencrypt import StandardEncryption

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select

from app.extensions import db
from app.models import Category, CompetitionDay, Gymnast, ScoreSummary, ScoreEntry, JudgeRole, SubmissionStatus
from app.services.championship_operations_service import assignments_effective_for_category
from app.services.cloud_scoring_service import refresh_score_summary


def _categories_data(championship):
    categories = db.session.execute(select(Category).where(
        Category.championship_id == championship.id,
        Category.deleted_at.is_(None),
    ).order_by(Category.competition_day_id, Category.bench, Category.session,
               Category.passing_order)).scalars().all()
    days = {
        day.id: day for day in db.session.execute(select(CompetitionDay).where(
            CompetitionDay.championship_id == championship.id
        )).scalars()
    }
    result = []
    for category in categories:
        gymnasts = db.session.execute(select(Gymnast).where(
            Gymnast.category_id == category.id,
            Gymnast.deleted_at.is_(None),
        ).order_by(Gymnast.passing_order, Gymnast.id)).scalars().all()
        summaries = {gymnast.id: refresh_score_summary(gymnast.id) for gymnast in gymnasts}
        ranked = sorted(gymnasts, key=lambda gymnast: (
            -summaries[gymnast.id].total_score,
            -summaries[gymnast.id].e_score,
            -summaries[gymnast.id].a_score,
            gymnast.passing_order,
        ))
        result.append({
            'category': category,
            'day': days[category.competition_day_id],
            'rows': [{
                'position': index,
                'gymnast': gymnast,
                'summary': summaries[gymnast.id],
            } for index, gymnast in enumerate(ranked, start=1)],
        })
    return result


def build_excel(championship):
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill('solid', fgColor='2D3561')
    header_font = Font(color='FFFFFF', bold=True)
    categories_data = _categories_data(championship)
    if not categories_data:
        worksheet = workbook.create_sheet('Resultados')
        worksheet.append(['Campeonato', championship.name])
    for category_data in categories_data:
        category = category_data['category']
        day = category_data['day']
        worksheet = workbook.create_sheet(category.name[:31] or 'Resultados')
        roles = [JudgeRole.DB, JudgeRole.DA, JudgeRole.A, JudgeRole.E]
        assignments = sorted(
            (assignment for assignment in assignments_effective_for_category(
                category, include_superseded=True
            ) if assignment.role in roles),
            key=lambda assignment: (
                roles.index(assignment.role), assignment.created_at, str(assignment.id)
            ),
        )
        area_columns = []
        score_headers = []
        for role in roles:
            peers = [a for a in assignments if a.role == role]
            area_columns.append((role, peers))
            score_headers.extend(
                f'{role.value}{index}' for index in range(1, len(peers) + 1)
            )
            score_headers.append(f'{role.value} total')
        gymnast_ids = [row['gymnast'].id for row in category_data['rows']]
        entries = db.session.execute(select(ScoreEntry).where(
            ScoreEntry.gymnast_id.in_(gymnast_ids),
            ScoreEntry.judge_assignment_id.in_([a.id for a in assignments]),
        )).scalars().unique().all() if gymnast_ids and assignments else []
        values = {
            (entry.gymnast_id, entry.judge_assignment_id): float(entry.value)
            for entry in entries if entry.submission_status == SubmissionStatus.SUBMITTED
        }
        worksheet.append([
            'Pos.', 'Nombre', 'Club', *score_headers,
            'Descuento', 'Total', 'Día', 'Banca', 'Jornada',
        ])
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
        for row in category_data['rows']:
            gymnast, summary = row['gymnast'], row['summary']
            scores = []
            for role, peers in area_columns:
                scores.extend(values.get((gymnast.id, a.id)) for a in peers)
                scores.append(float(getattr(summary, f'{role.value.lower()}_score')))
            worksheet.append([
                row['position'], gymnast.full_name, gymnast.club_name,
                *scores,
                float(summary.discount), float(summary.total_score),
                day.competition_date.isoformat(), category.bench.value,
                category.session.value,
            ])
        for column in worksheet.columns:
            worksheet.column_dimensions[column[0].column_letter].width = min(
                max(len(str(cell.value or '')) for cell in column) + 2, 32
            )
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_pdf(championship):
    output = BytesIO()
    document = SimpleDocTemplate(
        output, pagesize=landscape(A4), leftMargin=30, rightMargin=30,
        topMargin=30, bottomMargin=30,
        encrypt=StandardEncryption('', ownerPassword=token_urlsafe(32),
                                   canPrint=1, canModify=0, canCopy=0, canAnnotate=0),
    )
    styles = getSampleStyleSheet()
    story = [Paragraph(f'Resultados — {escape(championship.name)}', styles['Title'])]
    for category_data in _categories_data(championship):
        category = category_data['category']
        day = category_data['day']
        story.extend([
            Spacer(1, 10),
            Paragraph(
                f'{escape(category.name)} · Día {day.sequence} · Banca {category.bench.value} · {category.session.value}',
                styles['Heading3'],
            ),
        ])
        rows = [['Pos.', 'Nombre', 'Club', 'DB', 'DA', 'Desc.', 'Total']]
        for row in category_data['rows']:
            gymnast, summary = row['gymnast'], row['summary']
            rows.append([
                row['position'],
                Paragraph(escape(gymnast.full_name), styles['BodyText']),
                Paragraph(escape(gymnast.club_name or ''), styles['BodyText']),
                f'{summary.db_score:.2f}', f'{summary.da_score:.2f}',
                f'{summary.discount:.2f}', f'{summary.total_score:.2f}',
            ])
        table = Table(rows, colWidths=[32, 270, 220, 68, 68, 68, 80], repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2D3561')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#CBD5E1')),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F1F5F9')]),
            ('ALIGN', (1, 1), (2, -1), 'LEFT'),
        ]))
        story.append(table)
    document.build(story)
    output.seek(0)
    return output
