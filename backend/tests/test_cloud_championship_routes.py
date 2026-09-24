from io import BytesIO

import openpyxl
from sqlalchemy import select

from app.extensions import db
from app.models import (
    AccountType,
    AuditLog,
    Category,
    CompetitionDay,
    Gymnast,
    User,
)
from app.security.passwords import hash_password


def create_admin():
    admin = User(
        account_type=AccountType.GLOBAL_ADMIN,
        first_name='Administrador',
        last_name='Global',
        username='ADMIN',
        password_hash=hash_password('Clave-Admin-123'),
    )
    db.session.add(admin)
    db.session.commit()
    return admin


def login_headers(client):
    response = client.post(
        '/api/v1/auth/login',
        json={'username': 'ADMIN', 'password': 'Clave-Admin-123'},
    )
    assert response.status_code == 200
    csrf_cookie = client.get_cookie('ritmica_csrf')
    return {'X-CSRF-TOKEN': csrf_cookie.value}


def small_excel():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'SABADO'
    sheet.append([
        'N', 'NOMBRE', 'CLUB', 'CATEGORIA',
        'N', 'NOMBRE', 'CLUB', 'CATEGORIA',
    ])
    sheet.append([
        1, 'Gimnasta A1', 'Club A', 'MINI A',
        1, 'Gimnasta B1', 'Club B', 'MINI B',
    ])
    sheet.append([
        2, 'Gimnasta A2', 'Club A', 'MINI A',
        2, 'Gimnasta B2', 'Club B', 'MINI B',
    ])
    sheet.append([
        'BREAK 1 HORA - PREMIACION - ALMUERZO JUECES',
        None, None, None, None, None, None, None,
    ])
    sheet.append([
        'N', 'NOMBRE', 'CLUB', 'CATEGORIA',
        'N', 'NOMBRE', 'CLUB', 'CATEGORIA',
    ])
    sheet.append([
        1, 'Gimnasta A3', 'Club A', 'JUNIOR A',
        1, 'Gimnasta B3', 'Club B', 'JUNIOR B',
    ])
    judges = workbook.create_sheet("Jueces")
    judges.append(["JORNADA AM", None, None, None, None, None])
    judges.append(["María Pérez", "12.345.678-5", "DA", None, None, None])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output


def test_admin_imports_preview_and_confirms_order_of_passage(app, client):
    create_admin()
    headers = login_headers(client)

    create_response = client.post(
        '/api/v1/championships',
        json={
            'name': 'Clasificatorio Centro',
            'kind': 'clasificatorio',
            'qualifier_number': 1,
            'zone': 'centro',
            'start_date': '2026-08-01',
        },
        headers=headers,
    )
    assert create_response.status_code == 201
    championship_id = create_response.get_json()['championship']['id']

    upload_response = client.post(
        f'/api/v1/championships/{championship_id}/import-previews',
        data={'file': (small_excel(), 'orden-centro.xlsx'), 'competition_date': '2026-08-01'},
        headers=headers,
        content_type='multipart/form-data',
    )
    assert upload_response.status_code == 201
    preview_payload = upload_response.get_json()
    preview_id = preview_payload['id']
    sheet = preview_payload['preview']['sheets'][0]
    assert sheet['detected_cutoff_row'] == 4
    assert sheet['selected_cutoff_row'] == 4
    assert not sheet['cutoff_confirmed']
    assert preview_payload['preview']['total_categories'] == 4
    assert preview_payload['preview']['total_gymnasts'] == 6

    premature = client.post(
        f'/api/v1/championships/{championship_id}'
        f'/import-previews/{preview_id}/confirm',
        headers=headers,
    )
    assert premature.status_code == 400
    assert premature.get_json()['code'] == 'UNCONFIRMED_CUTOFF'

    update_response = client.patch(
        f'/api/v1/championships/{championship_id}'
        f'/import-previews/{preview_id}',
        json={'accept_detected': True},
        headers=headers,
    )
    assert update_response.status_code == 200
    assert (
        update_response.get_json()['preview']['sheets'][0]
        ['cutoff_confirmed']
    )

    confirmation = client.post(
        f'/api/v1/championships/{championship_id}'
        f'/import-previews/{preview_id}/confirm',
        headers=headers,
    )
    assert confirmation.status_code == 200
    imported = confirmation.get_json()['imported']
    assert {key: imported[key] for key in ('days', 'categories', 'gymnasts')} == {
        'days': 1,
        'categories': 4,
        'gymnasts': 6,
    }
    assert db.session.execute(select(CompetitionDay)).scalars().all()
    assert len(db.session.execute(select(Category)).scalars().all()) == 4
    assert len(db.session.execute(select(Gymnast)).scalars().all()) == 6

    detail = client.get(
        f'/api/v1/championships/{championship_id}',
    )
    assert detail.status_code == 200
    assert detail.get_json()['championship']['counts'] == {
        'days': 1,
        'categories': 4,
        'gymnasts': 6,
    }

    actions = db.session.execute(
        select(AuditLog.action).order_by(AuditLog.occurred_at)
    ).scalars().all()
    assert actions == [
        'LOGIN_SUCCESS',
        'CHAMPIONSHIP_CREATED',
        'IMPORT_PREVIEW_CREATED',
        'IMPORT_CUTOFFS_UPDATED',
        'IMPORT_CONFIRMED',
    ]


def test_cloud_championship_routes_require_login(client):
    response = client.get('/api/v1/championships')

    assert response.status_code == 401
    assert response.get_json()['code'] == 'AUTHENTICATION_REQUIRED'


def test_qualifier_number_is_required_and_persisted(app, client):
    create_admin()
    headers = login_headers(client)
    payload = {'name': 'Clasificatorio', 'kind': 'CLASIFICATORIO',
               'zone': 'NORTE', 'start_date': '2026-08-01'}
    for invalid in (None, 0, 3, True, '1'):
        response = client.post('/api/v1/championships',
                               json={**payload, 'qualifier_number': invalid}, headers=headers)
        assert response.status_code == 400
    for number in (1, 2):
        response = client.post('/api/v1/championships',
                               json={**payload, 'qualifier_number': number}, headers=headers)
        assert response.status_code == 201
        saved = response.get_json()['championship']
        assert saved['qualifier_number'] == number
        detail = client.get(f"/api/v1/championships/{saved['id']}", headers=headers)
        assert detail.get_json()['championship']['qualifier_number'] == number


def test_zone_and_final_validation(app, client):
    create_admin()
    headers = login_headers(client)
    payload = {'name': 'Final', 'kind': 'FINAL', 'zone': 'SUR', 'start_date': '2026-08-01'}
    assert client.post('/api/v1/championships', json={**payload, 'zone': 'OTRA'},
                       headers=headers).status_code == 400
    assert client.post('/api/v1/championships', json={**payload, 'qualifier_number': 1},
                       headers=headers).status_code == 400
    response = client.post('/api/v1/championships', json=payload, headers=headers)
    assert response.status_code == 201
    assert response.get_json()['championship']['qualifier_number'] is None


def test_validate_day_file_before_creating_championship(app, client):
    from app.models import Championship, FileObject, ImportPreview

    create_admin()
    headers = login_headers(client)
    response = client.post('/api/v1/championships/validate-day-file',
                           headers=headers, data={'file': (small_excel(), 'dia.xlsx')})
    assert response.status_code == 200
    assert response.json == {'valid': True}
    for model in (Championship, FileObject, ImportPreview):
        assert db.session.execute(select(model)).scalars().all() == []


def test_validate_day_file_reports_category_conflict(app, client):
    create_admin()
    headers = login_headers(client)
    workbook = openpyxl.load_workbook(small_excel())
    workbook['SABADO']['H2'] = 'MINI A'
    workbook['Jueces']['F1'] = 'BANCA B'
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    response = client.post('/api/v1/championships/validate-day-file',
                           headers=headers, data={'file': (output, 'dia.xlsx')})
    assert response.status_code == 400
    assert 'MINI A' in response.json['error']


def test_validate_day_file_requires_login(client):
    response = client.post('/api/v1/championships/validate-day-file',
                           data={'file': (small_excel(), 'dia.xlsx')})
    assert response.status_code == 401
